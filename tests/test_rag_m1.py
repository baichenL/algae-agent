from pathlib import Path

from app.core import database
from app.models.rag_schema import RagQueryRequest
from app.services.rag.ingestion.index_store import ingest_file
from app.services.rag.ingestion.ingest import ingest_sources
from app.services.rag.ingestion.metadata import parse_source_metadata
from app.services.rag.service import answer_rag_question


def test_parse_rag_filename():
    metadata = parse_source_metadata("media_recipe__tap_medium__v1__zh.docx")

    assert metadata["doc_type"] == "media_recipe"
    assert metadata["topic"] == "tap_medium"
    assert metadata["version"] == "v1"
    assert metadata["language"] == "zh"

    paper = parse_source_metadata("paper__microalgal_kinetics__2019__en.pdf")
    assert paper["doc_type"] == "paper"
    assert paper["year"] == "2019"


def test_ingest_docx_creates_document_and_chunks(isolated_sqlite_db, tmp_path):
    from docx import Document

    path = tmp_path / "media_recipe__tap_medium__v1__zh.docx"
    document = Document()
    document.add_heading("TAP medium", level=1)
    document.add_paragraph("TAP medium contains Tris, NH4Cl, MgSO4, CaCl2, phosphate buffer, and trace elements.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Tris"
    table.cell(1, 1).text = "2.42 g/L"
    document.save(path)

    result = ingest_file(path)

    assert result["status"] == "indexed"
    assert result["chunk_count"] >= 1
    docs = database.list_rag_documents()
    assert docs[0]["status"] == "indexed"
    assert docs[0]["doc_type"] == "media_recipe"


def test_ingest_xlsx_tracks_sheet_and_rows(isolated_sqlite_db, tmp_path):
    import openpyxl

    path = tmp_path / "experiment_data__algae_growth_curve__v1__zh.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "OD750"
    sheet.append(["strain_id", "day", "od750"])
    sheet.append(["Chlorella_01", 1, 0.2])
    sheet.append(["Chlorella_01", 2, 0.4])
    workbook.save(path)

    result = ingest_file(path)
    response = answer_rag_question(RagQueryRequest(question="Chlorella_01 OD750", top_k=3))

    assert result["status"] == "indexed"
    assert response.citations
    assert response.citations[0].sheet_name == "OD750"
    assert response.citations[0].row_start is not None


def test_ingest_skips_unchanged_file(isolated_sqlite_db, tmp_path):
    path = tmp_path / "manual__simple_sop__v1__en.txt"
    path.write_text("Sterilize workspace before microalgae transfer.", encoding="utf-8")

    first = ingest_file(path)
    second = ingest_file(path)

    assert first["status"] == "indexed"
    assert second["status"] == "skipped"
    assert database.list_rag_documents()[0]["status"] == "skipped"


def test_ingest_marks_failed_file(isolated_sqlite_db, tmp_path, monkeypatch):
    from app.services.rag.ingestion import index_store

    path = tmp_path / "manual__broken__v1__en.txt"
    path.write_text("content", encoding="utf-8")

    def fail_loader(path, metadata):
        raise RuntimeError("loader failed")

    monkeypatch.setattr(index_store, "load_source_units", fail_loader)
    result = ingest_file(path)

    assert result["status"] == "failed"
    assert "loader failed" in database.list_rag_documents()[0]["error_message"]


def test_rag_query_returns_citations_and_uncertainty(isolated_sqlite_db, tmp_path):
    path = tmp_path / "manual__contamination_check_sop__v1__en.txt"
    path.write_text("Contamination checks include microscopy, color inspection, and growth curve review.", encoding="utf-8")
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="How to check contamination?", top_k=2))

    assert response.status == "success"
    assert response.answer is not None
    assert response.citations
    assert response.answer.evidence
    assert response.answer.uncertainty


def test_rag_blocks_execution_requests(isolated_sqlite_db):
    response = answer_rag_question(RagQueryRequest(question="请执行传代并发送邮件"))

    assert response.blocked is True
    assert response.blocked_reason == "rag_read_only_boundary"
    assert response.answer is None


def test_doc_type_priority_prefers_media_recipe_over_paper(isolated_sqlite_db, tmp_path):
    recipe = tmp_path / "media_recipe__tap_medium__v1__zh.txt"
    paper = tmp_path / "paper__tap_medium_review__2024__en.txt"
    recipe.write_text("TAP medium recipe: Tris, NH4Cl, MgSO4, CaCl2, phosphate buffer, trace elements.", encoding="utf-8")
    paper.write_text("A paper mentions TAP medium and microalgae cultivation.", encoding="utf-8")
    ingest_sources(str(tmp_path))

    response = answer_rag_question(RagQueryRequest(question="TAP medium recipe", top_k=2))

    assert response.citations
    assert response.citations[0].doc_type == "media_recipe"


def test_hybrid_retrieval_finds_semantic_manual_match_without_exact_query_terms(isolated_sqlite_db, tmp_path):
    source = tmp_path / "manual__culture_quality_check__v1__en.txt"
    source.write_text(
        "Culture contamination checks include microscopy inspection, color review, and growth curve review.",
        encoding="utf-8",
    )
    ingest_file(source)

    response = answer_rag_question(RagQueryRequest(question="How do I detect a polluted culture?", top_k=3))

    assert response.answer is not None
    assert response.citations
    assert response.citations[0].file_name == source.name
    assert response.answer.facts
    assert response.answer.facts[0].citation_ids


def test_controlled_answer_segments_bind_facts_to_citations(isolated_sqlite_db, tmp_path):
    source = tmp_path / "manual__contamination_check_sop__v1__en.txt"
    source.write_text(
        "Contamination checks include microscopy, color inspection, and growth curve review.",
        encoding="utf-8",
    )
    ingest_file(source)

    response = answer_rag_question(RagQueryRequest(question="How to check contamination?", top_k=2))

    assert response.answer is not None
    assert response.answer.facts
    assert response.answer.explanations
    assert response.answer.suggestions
    assert response.answer.facts[0].citation_ids == [response.citations[0].source_id]
    assert "执行" in response.answer.uncertainty[-1] or "workflow" in response.answer.uncertainty[-1]


def test_m1_fallback_support_question_requires_condition_and_relation_coverage(isolated_sqlite_db, tmp_path):
    source = tmp_path / "media_recipe__tap_medium__v1__en.txt"
    source.write_text(
        "TAP medium recipe: Tris, NH4Cl, MgSO4, CaCl2, phosphate buffer, and trace elements.",
        encoding="utf-8",
    )
    ingest_file(source)

    response = answer_rag_question(
        RagQueryRequest(question="Does TAP documentation support anaerobic fermentation?", top_k=2)
    )

    assert response.answer is not None
    sufficiency = response.answer.debug["evidence_sufficiency"]
    assert sufficiency["status"] == "background_only"
    assert "condition" in sufficiency["missing_aspects"]
    assert "relation" in sufficiency["missing_aspects"]
    assert "no direct evidence" in response.answer.conclusion.lower()
    assert response.answer.facts
    assert response.answer.facts[0].text.startswith("背景资料：")
    assert "缺失内容" in response.answer.evidence[0].quote_summary
