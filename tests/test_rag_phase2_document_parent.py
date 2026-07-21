from pathlib import Path

from app.core import database
from app.services.rag.ingestion.index_store import ingest_file
from app.services.rag.ingestion.metadata import compute_file_hash, parse_source_metadata
from app.services.rag.ingestion.parsers import parse_document
from app.services.rag.retrieval.parent_expansion import expand_parent_context
from app.services.rag.retrieval.retriever import retrieve_chunks


def test_markdown_parser_preserves_heading_path(tmp_path):
    path = tmp_path / "manual__md_sop__v1__en.md"
    path.write_text("# Root\n\nIntro.\n\n## Step A\n\n- item one\n- item two\n", encoding="utf-8")
    metadata = {**parse_source_metadata(path), "content_hash": compute_file_hash(path)}

    document = parse_document(path, metadata)

    assert document.elements[0].element_type == "heading"
    assert any(element.section_path == ["Root", "Step A"] for element in document.elements)
    assert all(element.previous_id or index == 0 for index, element in enumerate(document.elements))
    assert all(element.next_id or index == len(document.elements) - 1 for index, element in enumerate(document.elements))


def test_docx_parser_preserves_heading_hierarchy(tmp_path):
    from docx import Document

    path = tmp_path / "manual__docx_sop__v1__en.docx"
    doc = Document()
    doc.add_heading("Main SOP", level=1)
    doc.add_paragraph("Main paragraph.")
    doc.add_heading("Recovery", level=2)
    doc.add_paragraph("Recovery step.")
    doc.save(path)
    metadata = {**parse_source_metadata(path), "content_hash": compute_file_hash(path)}

    document = parse_document(path, metadata)

    paragraph = next(element for element in document.elements if element.content == "Recovery step.")
    assert paragraph.section_path == ["Main SOP", "Recovery"]
    assert paragraph.parent_id


def test_xlsx_parser_builds_table_parent_and_row_children(tmp_path):
    import openpyxl

    path = tmp_path / "experiment_data__growth__v1__en.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Growth"
    sheet.append(["time h", "biomass g/L"])
    sheet.append([0, 0.1])
    workbook.save(path)
    metadata = {**parse_source_metadata(path), "content_hash": compute_file_hash(path)}

    document = parse_document(path, metadata)

    table = next(element for element in document.elements if element.element_type == "table")
    row = next(element for element in document.elements if element.element_type == "table_row")
    assert table.source_locator["sheet_name"] == "Growth"
    assert row.parent_id == table.element_id
    assert row.source_locator["row_index"] == 2


def test_ingest_persists_document_elements_without_duplicates(isolated_sqlite_db, tmp_path):
    path = tmp_path / "manual__simple_elements__v1__en.md"
    path.write_text("# SOP\n\nFirst paragraph.\n\nSecond paragraph.", encoding="utf-8")

    first = ingest_file(path)
    first_elements = database.list_rag_document_elements()
    second = ingest_file(path, rebuild=True)
    second_elements = database.list_rag_document_elements()

    assert first["status"] == "indexed"
    assert second["status"] == "indexed"
    assert len(first_elements) == len(second_elements)
    assert {item["element_id"] for item in first_elements} == {item["element_id"] for item in second_elements}


def test_parent_expansion_adds_context_and_can_be_disabled(isolated_sqlite_db, tmp_path, monkeypatch):
    path = tmp_path / "manual__parent_context__v1__en.md"
    path.write_text("# SOP\n\nStep one mentions microscopy.\n\nStep two mentions color inspection.", encoding="utf-8")
    ingest_file(path)

    chunks = retrieve_chunks("microscopy inspection", top_k=1, doc_types=["manual"])
    assert chunks
    assert chunks[0]["parent_expansion"]["enabled"] is True
    assert "answer_context" in chunks[0]

    monkeypatch.setenv("RAG_PARENT_EXPANSION_ENABLED", "false")
    disabled = expand_parent_context(chunks)
    assert disabled[0]["parent_expansion"]["enabled"] is False


def test_docling_flag_falls_back_without_dependency(tmp_path, monkeypatch):
    path = tmp_path / "manual__docling_fallback__v1__en.txt"
    path.write_text("Fallback parser content.", encoding="utf-8")
    metadata = {**parse_source_metadata(path), "content_hash": compute_file_hash(path)}
    monkeypatch.setenv("RAG_DOCLING_ENABLED", "true")

    document = parse_document(path, metadata)

    assert document.parser_name == "fallback_structured_parser"
    assert document.metadata["parser_warning"] == "docling_requested_but_not_configured_fallback_used"
    assert document.elements
