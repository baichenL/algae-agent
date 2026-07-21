from pathlib import Path

import openpyxl
from docx import Document

from app.core import database
from app.models.rag_schema import RagQueryRequest
from app.models.rag_evidence_schema import EvidenceUnit, QueryFrame
from app.models.rag_semantic_schema import EvidenceContract, SemanticSubquery, TypedEntity
from app.services.rag.ingestion.index_store import ingest_file
from app.services.rag.service import answer_rag_question
from app.services.rag.evidence.adapters.chunk_adapter import chunk_to_evidence_unit
from app.services.rag.evidence.adapters.table_schema_adapter import (
    extract_table_schema_rows,
    infer_column_role,
    normalize_column_name,
)
from app.services.rag.evidence.backfill_schema import backfill_table_schemas
from app.services.rag.evidence.query_frame import parse_query_frame
from app.services.rag.evidence.semantic_query import semantic_subquery_to_frame


def _make_growth_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append([None, "time h", "light intensity", "N mg/L", "P mg/L", "tf", "biomass g/L"])
    sheet.append([0, 0, 26, 0.335, 0.0157, 12, 0.1138244])
    sheet.append([1, 12, 26, 0.335, 0.0157, 12, 0.1182614])
    workbook.save(path)


def _make_tap_recipe_docx(path: Path) -> None:
    document = Document()
    document.add_heading("TAP medium recipe", level=1)

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "component"
    table.rows[0].cells[1].text = "amount /g"
    for name, amount in [("NH4Cl", "20.0"), ("MgSO4 7H2O", "5.0"), ("CaCl2 2H2O", "2.5")]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = amount

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "component"
    table.rows[0].cells[1].text = "amount /g"
    for name, amount in [("K2HPO4", "10.8"), ("KH2PO4", "5.6")]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = amount

    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "component"
    table.rows[0].cells[1].text = "amount /g"
    table.rows[0].cells[2].text = "H2O /ml"
    for name, amount, water in [
        ("Na2-EDTA2H2O", "50", "250"),
        ("ZnSO4 7H2O", "22", "100"),
        ("H3BO3", "11.4", "200"),
        ("FeSO4 7H2O", "4.99", "50"),
    ]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = amount
        cells[2].text = water

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "component"
    table.rows[0].cells[1].text = "amount"
    for name, amount in [
        ("Tris(sigma)", "2.42g"),
        ("salt stock solution", "10ml"),
        ("phosphate stock solution", "1 mL"),
        ("Hutner trace elements", "1 mL"),
        ("acetic acid", "1 mL"),
        ("final volume", "1 L"),
    ]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = amount

    document.save(path)


def _make_manual_txt(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "Electroporation protocol for Chlamydomonas.",
                "Grow cells to OD750=0.3-0.5 before electroporation.",
                "Use 40 mM sucrose TAP medium for resuspension and recovery culture.",
                "Hygromycin TAP plates are used for post-transformation screening.",
            ]
        ),
        encoding="utf-8",
    )


def _make_paper_txt(path: Path) -> None:
    path.write_text(
        "This review discusses machine learning, data-driven modeling, neural network methods "
        "and growth prediction for microalgae cultivation systems. It is a review source, not a lab SOP.",
        encoding="utf-8",
    )


def test_rag_evidence_schema_models_can_represent_closed_world_negative():
    frame = QueryFrame(
        original_question="experiment data has pH column?",
        question_type="existence",
        target_entity="experiment_data",
        target_attribute="column",
        target_value="pH",
        answer_shape="yes_no",
        evidence_requirement="structured_schema",
        allow_inference=False,
        source_constraint=["experiment_data"],
    )
    evidence = EvidenceUnit(
        evidence_id="schema:file:sheet:1:n_mg_l",
        source_type="experiment_data",
        source_file="file.xlsx",
        fact_type="table_column",
        entity="file.xlsx:Sheet1",
        attribute="column_name",
        value="N mg/L",
        location={"sheet_name": "Sheet1", "column_index": 1},
    )

    assert frame.target_value == "pH"
    assert evidence.fact_type == "table_column"


def test_chunk_adapter_preserves_location_and_open_world_type():
    unit = chunk_to_evidence_unit(
        {
            "chunk_id": "abc",
            "doc_type": "manual",
            "file_name": "manual.pdf",
            "source_path": "data/manual.pdf",
            "title": "manual",
            "page_number": 12,
            "chunk_index": 3,
            "content": "A manual text span.",
            "metadata": {"topic": "manual"},
        }
    )

    assert unit.fact_type == "text_chunk"
    assert unit.location["page_number"] == 12
    assert unit.citation["chunk_id"] == "abc"


def test_table_schema_adapter_extracts_true_columns(tmp_path):
    path = tmp_path / "experiment_data__algae_growth_curve__v1__en.xlsx"
    _make_growth_xlsx(path)

    rows = extract_table_schema_rows(
        path,
        {"doc_type": "experiment_data", "source_path": str(path)},
        document_id=1,
    )
    column_names = [row["column_name"] for row in rows]

    assert "light intensity" in column_names
    assert "N mg/L" in column_names
    assert "P mg/L" in column_names
    assert "tf" in column_names
    assert "biomass g/L" in column_names
    assert "pH" not in column_names
    assert normalize_column_name("N mg/L") == "n_mg_l"
    assert infer_column_role("biomass g/L") == "biomass"


def test_ingest_xlsx_persists_table_schema(isolated_sqlite_db, tmp_path):
    path = tmp_path / "experiment_data__algae_growth_curve__v1__en.xlsx"
    _make_growth_xlsx(path)

    result = ingest_file(path)
    schemas = database.list_rag_source_schemas(doc_types=["experiment_data"])

    assert result["status"] == "indexed"
    assert {row["column_name"] for row in schemas} >= {
        "time h",
        "light intensity",
        "N mg/L",
        "P mg/L",
        "tf",
        "biomass g/L",
    }


def test_query_frame_parser_identifies_environment_variable_list():
    frame = parse_query_frame("experiment data environmental variable list")

    assert frame.question_type == "list"
    assert frame.target_attribute == "environment_variables"
    assert frame.evidence_requirement == "structured_schema"


def test_m2_lists_only_true_environment_variable_columns(isolated_sqlite_db, tmp_path):
    path = tmp_path / "experiment_data__algae_growth_curve__v1__en.xlsx"
    _make_growth_xlsx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="experiment data environmental variable list"))

    assert response.answer is not None
    conclusion = response.answer.conclusion
    assert "light intensity" in conclusion
    assert "N mg/L" in conclusion
    assert "P mg/L" in conclusion
    assert "tf" in conclusion
    assert "biomass g/L" not in conclusion
    assert "pH" not in conclusion


def test_m2_data_analysis_answers_max_biomass(isolated_sqlite_db, tmp_path):
    path = tmp_path / "experiment_data__algae_growth_curve__v1__en.xlsx"
    _make_growth_xlsx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="experiment data max biomass"))

    assert response.answer is not None
    assert "0.1182614" in response.answer.conclusion
    assert response.citations
    assert response.answer.facts
    assert response.answer.facts[0].citation_ids


def test_m2_data_analysis_refuses_unsupported_relationship(isolated_sqlite_db, tmp_path):
    path = tmp_path / "experiment_data__algae_growth_curve__v1__en.xlsx"
    _make_growth_xlsx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="experiment data pH effect on biomass correlation"))

    assert response.answer is not None
    assert "pH" in response.answer.conclusion
    assert response.debug["answerability"]["status"] in {"not_found", "partial"}


def test_ingest_docx_persists_recipe_components(isolated_sqlite_db, tmp_path):
    path = tmp_path / "media_recipe__tap_medium__v1__en.docx"
    _make_tap_recipe_docx(path)

    ingest_file(path)
    rows = database.list_rag_recipe_components(doc_types=["media_recipe"])

    assert {row["component_name"] for row in rows} >= {"NH4Cl", "K2HPO4", "KH2PO4", "Tris(sigma)"}
    assert any(row["group_name"] == "phosphate_solution" for row in rows)


def test_recipe_adapter_answers_component_amount_with_citations(isolated_sqlite_db, tmp_path):
    path = tmp_path / "media_recipe__tap_medium__v1__en.docx"
    _make_tap_recipe_docx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="TAP recipe K2HPO4 amount"))

    assert response.answer is not None
    assert "10.8" in response.answer.conclusion
    assert response.citations
    assert response.answer.facts
    assert set(response.answer.facts[0].citation_ids).issubset({item.source_id for item in response.citations})


def test_recipe_adapter_closed_world_absence_does_not_hallucinate_component(isolated_sqlite_db, tmp_path):
    path = tmp_path / "media_recipe__tap_medium__v1__en.docx"
    _make_tap_recipe_docx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="TAP recipe has NaCl component?"))

    assert response.answer is not None
    assert "NaCl" in response.answer.conclusion
    assert "NH4Cl" in response.answer.conclusion or "NH" in response.answer.conclusion
    assert response.debug["answerability"]["closed_world_negative"] is True


def test_waterbody_microalgae_question_does_not_default_to_tap_recipe(isolated_sqlite_db, tmp_path):
    path = tmp_path / "media_recipe__tap_medium__v1__zh.docx"
    _make_tap_recipe_docx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="沈阳的浑河里有哪些常见的微藻"))

    assert response.answer is not None
    assert "当前知识库没有检索到与该问题直接相关的可引用资料" in response.answer.conclusion
    assert "TAP" not in response.answer.conclusion
    assert response.citations == []


def test_unmentioned_tap_semantic_overview_is_rejected():
    subquery = SemanticSubquery(
        subquery_id="sq1",
        original_text="沈阳的浑河里有哪些常见的微藻",
        intent="overview_of",
        relation="overview_of",
        entities=[
            TypedEntity(
                entity_type="culture_medium",
                surface="TAP medium",
                canonical_id="TAP medium",
                role="subject",
            )
        ],
        evidence_contract=EvidenceContract(relation="overview_of"),
        confidence=0.95,
    )

    assert semantic_subquery_to_frame(subquery) is None


def test_sop_adapter_answers_electroporation_without_full_manual_summary(isolated_sqlite_db, tmp_path):
    path = tmp_path / "manual__microalgae_laboratory_manual__v1__en.txt"
    _make_manual_txt(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="manual electroporation OD750 procedure"))

    assert response.answer is not None
    assert "OD750" in response.answer.conclusion
    assert response.citations
    assert {item.doc_type for item in response.citations} == {"manual"}


def test_paper_adapter_summarizes_methods_by_paper(isolated_sqlite_db, tmp_path):
    path = tmp_path / "paper__machine_learning_approaches_microalgae_cultivation_systems__2024__en.txt"
    _make_paper_txt(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="paper machine learning methods for microalgae cultivation"))

    assert response.answer is not None
    assert "machine learning approaches microalgae cultivation systems (2024)" in response.answer.conclusion
    assert "machine learning" in response.answer.conclusion
    assert response.citations


def test_paper_adapter_refuses_project_parameters_without_direct_evidence(isolated_sqlite_db, tmp_path):
    path = tmp_path / "paper__machine_learning_approaches_microalgae_cultivation_systems__2024__en.txt"
    _make_paper_txt(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="paper directly applicable project culture parameter"))

    assert response.answer is not None
    assert "SOP" in " ".join(response.answer.uncertainty)
    assert response.debug["answerability"]["status"] == "not_found"


def test_sufficiency_gate_does_not_treat_background_tap_as_dark_culture_support(isolated_sqlite_db, tmp_path):
    path = tmp_path / "manual__tap_recovery_background__v1__en.txt"
    path.write_text(
        "\n".join(
            [
                "TAP medium recipe and recovery notes.",
                "Use 40 mM sucrose TAP medium for recovery culture after electroporation.",
                "Keep transformed plates in the dark for marker handling, but no statement links TAP suitability to dark culture.",
            ]
        ),
        encoding="utf-8",
    )
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="Is TAP medium suitable for dark culture?"))

    assert response.answer is not None
    sufficiency = response.answer.debug["evidence_sufficiency"]
    assert sufficiency["status"] == "background_only"
    assert "relation" in sufficiency["missing_aspects"]
    assert "no direct evidence" in response.answer.conclusion.lower()
    assert response.answer.facts
    assert response.answer.facts[0].text.startswith("背景资料：")
    assert all("缺失内容" in item.quote_summary for item in response.answer.evidence)
    assert all("page 10" not in item.location for item in response.answer.evidence)


def test_sufficiency_gate_allows_explicit_dark_culture_statement(isolated_sqlite_db, tmp_path):
    path = tmp_path / "manual__tap_dark_culture_direct__v1__en.txt"
    path.write_text(
        "TAP medium is suitable for dark culture of Chlamydomonas under this protocol.",
        encoding="utf-8",
    )
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="Is TAP medium suitable for dark culture?"))

    assert response.answer is not None
    sufficiency = response.answer.debug["evidence_sufficiency"]
    assert sufficiency["status"] == "sufficient"
    assert sufficiency["direct_evidence_ids"]
    assert response.answer.facts
    assert response.answer.facts[0].citation_ids == [response.citations[0].source_id]


def test_m2_response_exposes_evidence_selection_debug(isolated_sqlite_db, tmp_path):
    path = tmp_path / "experiment_data__algae_growth_curve__v1__en.xlsx"
    _make_growth_xlsx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="experiment data max biomass"))

    assert response.debug["query_frame"]["target_entity"] == "experiment_data"
    assert response.debug["answerability"]["status"] == "answered"
    assert response.debug["evidence_selection"]["selected_count"] >= 1
    assert response.debug["evidence_selection"]["coverage"] == "sufficient"
    assert response.answer.explanations
    assert response.answer.suggestions


def test_m2_controlled_answer_keeps_safety_suggestion(isolated_sqlite_db, tmp_path):
    path = tmp_path / "media_recipe__tap_medium__v1__en.docx"
    _make_tap_recipe_docx(path)
    ingest_file(path)

    response = answer_rag_question(RagQueryRequest(question="TAP recipe NH4Cl amount"))

    assert response.answer is not None
    assert response.answer.facts
    assert response.answer.explanations
    assert response.answer.suggestions
    assert any("不会自动" in item.text or "审批" in item.text for item in response.answer.suggestions)


def test_schema_backfill_populates_existing_indexed_document(isolated_sqlite_db, tmp_path):
    path = tmp_path / "experiment_data__algae_growth_curve__v1__en.xlsx"
    _make_growth_xlsx(path)
    ingest_file(path)

    import sqlite3

    with sqlite3.connect(isolated_sqlite_db) as conn:
        conn.execute("DELETE FROM rag_source_schemas")
        conn.commit()

    assert database.list_rag_source_schemas(doc_types=["experiment_data"]) == []
    before_status = database.get_rag_index_status()["schema_status"]
    assert before_status["schema_count"] == 0
    assert before_status["documents_missing_schema"]

    summary = backfill_table_schemas(tmp_path)
    schemas = database.list_rag_source_schemas(doc_types=["experiment_data"])
    after_status = database.get_rag_index_status()["schema_status"]

    assert summary["counts"] == {"backfilled": 1}
    assert len(schemas) >= 5
    assert after_status["schema_count"] == len(schemas)
    assert after_status["documents_missing_schema"] == []
