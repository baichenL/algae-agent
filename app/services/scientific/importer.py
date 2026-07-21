from __future__ import annotations

import csv
import hashlib
import io
import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any


MAX_DATASET_BYTES = 10 * 1024 * 1024
TIME_ALIASES = {"time", "time_h", "hours", "hour", "elapsed_hours", "时间", "时间h", "培养时间"}
RESPONSE_ALIASES = {
    "biomass", "biomassg", "biomass_g_l", "od", "od600", "od750", "dry_weight", "growth", "生物量", "光密度"
}
IGNORED_ALIASES = {"", "id", "idx", "index", "row", "序号", "unnamed"}


class DatasetImportError(ValueError):
    def __init__(self, code: str, message: str, *, details: dict[str, Any] | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.details = details or {}


@dataclass(frozen=True)
class ParsedScientificDataset:
    dataset: dict[str, Any]
    batches: list[dict[str, Any]]


def _norm(value: Any) -> str:
    return "".join(str(value or "").strip().casefold().replace("（", "(").replace("）", ")").split())


def _alias_key(value: Any) -> str:
    text = _norm(value)
    for token in ("(", "[", "/"):
        text = text.split(token, 1)[0]
    return text.replace("·", "").replace("_", "")


def _numeric(value: Any, *, column: str, row: int) -> float:
    if value is None or value == "":
        raise DatasetImportError("missing_numeric_value", f"Missing numeric value in {column} at row {row}.")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise DatasetImportError(
            "non_numeric_value", f"Non-numeric value in {column} at row {row}.",
            details={"column": column, "row": row, "value": str(value)},
        ) from exc
    if not math.isfinite(number):
        raise DatasetImportError("non_finite_value", f"Non-finite value in {column} at row {row}.")
    return number


def _read_csv(content: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("gb18030")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise DatasetImportError("empty_dataset", "The uploaded CSV is empty.")
    return [str(item or "").strip() for item in rows[0]], rows[1:]


def _read_xlsx(content: bytes, sheet_name: str | None) -> tuple[list[str], list[list[Any]]]:
    try:
        import openpyxl
    except Exception as exc:
        raise DatasetImportError("dependency_missing", "openpyxl is required for XLSX import.") from exc
    formula_book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        formula_sheet = formula_book[sheet_name] if sheet_name else formula_book.active
        for row in formula_sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    raise DatasetImportError("formula_not_allowed", "XLSX formulas are not accepted as scientific observations.")
    finally:
        formula_book.close()
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = book[sheet_name] if sheet_name else book.active
        raw = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    if not raw:
        raise DatasetImportError("empty_dataset", "The uploaded workbook is empty.")
    return [str(item or "").strip() for item in raw[0]], [list(row) for row in raw[1:]]


def inspect_columns(headers: list[str]) -> dict[str, Any]:
    normalized = {_alias_key(header): header for header in headers}
    time_candidates = [original for key, original in normalized.items() if key in {_alias_key(x) for x in TIME_ALIASES}]
    response_candidates = [original for key, original in normalized.items() if key in {_alias_key(x) for x in RESPONSE_ALIASES}]
    return {
        "columns": headers,
        "time_candidates": time_candidates,
        "response_candidates": response_candidates,
        "suggested_mapping": {
            "time_column": time_candidates[0] if len(time_candidates) == 1 else None,
            "response_column": response_candidates[0] if len(response_candidates) == 1 else None,
        },
    }


def parse_scientific_dataset(
    *,
    content: bytes,
    filename: str,
    dataset_name: str | None = None,
    strain_id: str | None = None,
    mapping: dict[str, Any] | None = None,
    sheet_name: str | None = None,
) -> ParsedScientificDataset:
    if not content:
        raise DatasetImportError("empty_dataset", "The uploaded dataset is empty.")
    if len(content) > MAX_DATASET_BYTES:
        raise DatasetImportError("dataset_too_large", "The dataset exceeds the 10 MB limit.")
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        headers, rows = _read_csv(content)
    elif suffix == ".xlsx":
        headers, rows = _read_xlsx(content, sheet_name)
    else:
        raise DatasetImportError("unsupported_extension", "Only CSV and XLSX scientific datasets are supported.")
    if len(set(headers)) != len(headers):
        raise DatasetImportError("duplicate_columns", "Dataset column names must be unique.", details={"columns": headers})

    suggestions = inspect_columns(headers)
    mapping = dict(mapping or {})
    time_column = mapping.get("time_column") or suggestions["suggested_mapping"]["time_column"]
    response_column = mapping.get("response_column") or suggestions["suggested_mapping"]["response_column"]
    if not time_column or not response_column or time_column not in headers or response_column not in headers:
        raise DatasetImportError(
            "mapping_required",
            "A unique time column and response column are required.",
            details=suggestions,
        )

    batch_column = mapping.get("batch_column")
    replicate_column = mapping.get("replicate_column")
    ignored = set(mapping.get("ignore_columns") or [])
    factor_columns = list(mapping.get("factor_columns") or [])
    if not factor_columns:
        reserved = {time_column, response_column, batch_column, replicate_column, *ignored}
        factor_columns = [
            header for header in headers
            if header not in reserved and _alias_key(header) not in {_alias_key(x) for x in IGNORED_ALIASES}
        ]
    for column in [*factor_columns, batch_column, replicate_column]:
        if column and column not in headers:
            raise DatasetImportError("unknown_mapping_column", f"Unknown mapped column: {column}")

    positions = {header: index for index, header in enumerate(headers)}
    grouped: dict[str, dict[str, Any]] = {}
    valid_rows = 0
    metric_name = str(mapping.get("metric_name") or response_column).strip().casefold().replace(" ", "_")
    unit = mapping.get("response_unit")
    for source_row, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        padded = list(row) + [None] * max(0, len(headers) - len(row))
        elapsed = _numeric(padded[positions[time_column]], column=time_column, row=source_row)
        response = _numeric(padded[positions[response_column]], column=response_column, row=source_row)
        condition = {
            column: _numeric(padded[positions[column]], column=column, row=source_row)
            for column in factor_columns
        }
        if batch_column:
            raw_batch = str(padded[positions[batch_column]] or "").strip()
            if not raw_batch:
                raise DatasetImportError("missing_batch_id", f"Missing batch id at row {source_row}.")
            grouping_key = raw_batch
        else:
            grouping_key = json.dumps(condition, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        batch_digest = hashlib.sha256(grouping_key.encode("utf-8")).hexdigest()[:12]
        replicate = 1
        if replicate_column:
            replicate = int(_numeric(padded[positions[replicate_column]], column=replicate_column, row=source_row))
        batch = grouped.setdefault(
            grouping_key,
            {"condition": condition, "measurements": [], "batch_digest": batch_digest},
        )
        batch["measurements"].append(
            {
                "elapsed_hours": elapsed,
                "metric_name": metric_name,
                "value": response,
                "unit": unit,
                "replicate_index": replicate,
                "source_row": source_row,
            }
        )
        valid_rows += 1
    if not grouped:
        raise DatasetImportError("no_valid_rows", "No valid scientific observations were found.")

    content_hash = hashlib.sha256(content).hexdigest()
    dataset_id = f"ds_{content_hash[:16]}"
    batches = []
    for item in grouped.values():
        batches.append(
            {
                "id": f"{dataset_id}_b_{item['batch_digest']}",
                "strain_id": strain_id,
                "condition": item["condition"],
                "measurements": sorted(item["measurements"], key=lambda x: (x["elapsed_hours"], x["replicate_index"])),
                "provenance": "real_import",
            }
        )
    normalized_mapping = {
        "time_column": time_column,
        "response_column": response_column,
        "batch_column": batch_column,
        "replicate_column": replicate_column,
        "factor_columns": factor_columns,
        "metric_name": metric_name,
        "response_unit": unit,
        "sheet_name": sheet_name,
    }
    return ParsedScientificDataset(
        dataset={
            "id": dataset_id,
            "name": dataset_name or Path(filename).stem,
            "strain_id": strain_id,
            "source_file": Path(filename).name,
            "content_hash": content_hash,
            "mapping": normalized_mapping,
            "row_count": valid_rows,
            "status": "ready",
        },
        batches=batches,
    )
