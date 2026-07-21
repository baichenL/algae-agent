from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass
class LoadedUnit:
    content: str
    title: str | None = None
    section: str | None = None
    page_number: int | None = None
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


def load_source_units(path: str | Path, metadata: dict) -> list[LoadedUnit]:
    source = Path(path)
    extension = source.suffix.lower()
    if extension == ".docx":
        return _load_docx(source, metadata)
    if extension == ".pdf":
        return _load_pdf(source, metadata)
    if extension in {".xlsx", ".xls"}:
        return _load_xlsx(source, metadata)
    if extension in {".md", ".txt", ".csv"}:
        return _load_text(source, metadata)
    raise ValueError(f"Unsupported RAG source extension: {extension}")


def _load_docx(path: Path, metadata: dict) -> list[LoadedUnit]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to ingest DOCX files") from exc

    document = Document(str(path))
    units: list[LoadedUnit] = []
    current_section = metadata.get("topic")
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        text = "\n".join(item for item in buffer if item.strip()).strip()
        if text:
            units.append(LoadedUnit(content=text, section=current_section, title=metadata.get("topic")))
        buffer = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = getattr(paragraph.style, "name", "") or ""
        if style_name.lower().startswith("heading"):
            flush()
            current_section = text
        else:
            buffer.append(text)
    flush()

    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            units.append(
                LoadedUnit(
                    content="\n".join(rows),
                    section=f"{current_section or metadata.get('topic')} table {table_index}",
                    title=metadata.get("topic"),
                    metadata={"table_index": table_index},
                )
            )
    return units


def _load_pdf(path: Path, metadata: dict) -> list[LoadedUnit]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required to ingest PDF files") from exc

    reader = PdfReader(str(path))
    title = metadata.get("topic")
    try:
        pdf_title = (reader.metadata or {}).get("/Title")
        if pdf_title:
            title = str(pdf_title)
    except Exception:
        pass

    units = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            units.append(
                LoadedUnit(
                    content=text,
                    title=title,
                    section=f"page {page_index}",
                    page_number=page_index,
                )
            )
    return units


def _load_xlsx(path: Path, metadata: dict) -> list[LoadedUnit]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to ingest XLSX files") from exc

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    units: list[LoadedUnit] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header_index = _find_header_index(rows)
        headers = [_clean_cell(value) for value in rows[header_index]]
        data_rows = rows[header_index + 1 :]
        batch: list[str] = []
        batch_start = header_index + 2
        for offset, row in enumerate(data_rows, start=header_index + 2):
            values = [_clean_cell(value) for value in row]
            if not any(values):
                continue
            pairs = []
            for header, value in zip(headers, values):
                if header or value:
                    pairs.append(f"{header or 'column'}={value}")
            batch.append("; ".join(pairs))
            if len(batch) >= 20:
                units.append(
                    LoadedUnit(
                        content="\n".join(batch),
                        title=metadata.get("topic"),
                        section=sheet.title,
                        sheet_name=sheet.title,
                        row_start=batch_start,
                        row_end=offset,
                    )
                )
                batch = []
                batch_start = offset + 1
        if batch:
            units.append(
                LoadedUnit(
                    content="\n".join(batch),
                    title=metadata.get("topic"),
                    section=sheet.title,
                    sheet_name=sheet.title,
                    row_start=batch_start,
                    row_end=batch_start + len(batch) - 1,
                )
            )
    workbook.close()
    return units


def _load_text(path: Path, metadata: dict) -> list[LoadedUnit]:
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    return [LoadedUnit(content=text, title=metadata.get("topic"), section=path.name)] if text else []


def _find_header_index(rows: list[tuple]) -> int:
    for index, row in enumerate(rows[:10]):
        values = [_clean_cell(value) for value in row]
        if sum(1 for value in values if value) >= 2:
            return index
    return 0


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
