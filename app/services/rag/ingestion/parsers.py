from __future__ import annotations

import csv
import os
import re
from pathlib import Path
from typing import Any

from app.services.rag.ingestion.document_models import (
    DocumentElement,
    NormalizedDocument,
    stable_content_hash,
    stable_document_id,
    stable_element_id,
)
from app.services.rag.ingestion.loaders import LoadedUnit, load_source_units
from app.services.rag.ingestion.pdf_router import parse_pdf_document


FALLBACK_PARSER_NAME = "fallback_structured_parser"
FALLBACK_PARSER_VERSION = "2026-07-phase2-v1"


def parse_document(source_path: str | Path, metadata: dict) -> NormalizedDocument:
    if _docling_requested():
        # Docling remains optional. The first implementation keeps the lightweight
        # local parser as the safe Windows-friendly fallback.
        metadata = {**metadata, "parser_warning": "docling_requested_but_not_configured_fallback_used"}
    source = Path(source_path)
    extension = source.suffix.lower()
    if extension == ".docx":
        return _parse_docx(source, metadata)
    if extension == ".md":
        return _parse_markdown(source, metadata)
    if extension in {".xlsx", ".xls"}:
        return _parse_xlsx(source, metadata)
    if extension == ".csv":
        return _parse_csv(source, metadata)
    if extension == ".pdf":
        return parse_pdf_document(source, metadata)
    return _parse_loaded_units(source, metadata)


def normalized_document_to_loaded_units(document: NormalizedDocument) -> list[LoadedUnit]:
    units = []
    for element in document.elements:
        if element.element_type == "heading":
            continue
        locator = element.source_locator or {}
        units.append(
            LoadedUnit(
                content=element.content,
                title=document.title,
                section=" / ".join(element.section_path) if element.section_path else document.title,
                page_number=element.page_number,
                sheet_name=locator.get("sheet_name"),
                row_start=locator.get("row_start") or locator.get("row_index"),
                row_end=locator.get("row_end") or locator.get("row_index"),
                metadata={
                    **element.metadata,
                    "element_id": element.element_id,
                    "element_type": element.element_type,
                    "parent_id": element.parent_id,
                    "previous_id": element.previous_id,
                    "next_id": element.next_id,
                    "section_path": element.section_path,
                    "source_locator": element.source_locator,
                    "parser_name": document.parser_name,
                    "parser_version": document.parser_version,
                    "document_id": document.document_id,
                    "document_version": document.document_version,
                },
            )
        )
    return units


def _base_document(source: Path, metadata: dict, elements: list[DocumentElement]) -> NormalizedDocument:
    document_id = stable_document_id(metadata["source_path"])
    content_hash = metadata.get("content_hash") or ""
    version = metadata.get("version") or metadata.get("year") or content_hash[:16] or "unversioned"
    return NormalizedDocument(
        document_id=document_id,
        document_version=version,
        source_path=metadata["source_path"],
        source_type="local_file",
        title=metadata.get("topic") or source.stem,
        elements=_link_elements(elements),
        metadata=dict(metadata),
        parser_name=FALLBACK_PARSER_NAME,
        parser_version=FALLBACK_PARSER_VERSION,
        content_hash=content_hash,
    )


def _make_element(
    *,
    source: Path,
    metadata: dict,
    index: int,
    element_type: str,
    content: str,
    section_path: list[str] | None = None,
    parent_id: str | None = None,
    page_number: int | None = None,
    source_locator: dict[str, Any] | None = None,
    element_metadata: dict[str, Any] | None = None,
) -> DocumentElement:
    document_id = stable_document_id(metadata["source_path"])
    locator = {
        "source_path": metadata["source_path"],
        "file_name": source.name,
        **(source_locator or {}),
    }
    clean_content = _normalize_text(content)
    return DocumentElement(
        element_id=stable_element_id(
            document_id=document_id,
            element_type=element_type,
            index=index,
            content=clean_content,
            source_locator=locator,
        ),
        document_id=document_id,
        element_type=element_type,
        content=clean_content,
        page_number=page_number,
        section_path=section_path or [],
        parent_id=parent_id,
        source_locator=locator,
        metadata=element_metadata or {},
        content_hash=stable_content_hash(clean_content, locator),
    )


def _link_elements(elements: list[DocumentElement]) -> list[DocumentElement]:
    linked = []
    previous_id = None
    for index, element in enumerate(elements):
        next_id = elements[index + 1].element_id if index + 1 < len(elements) else None
        linked.append(element.model_copy(update={"previous_id": previous_id, "next_id": next_id}))
        previous_id = element.element_id
    return linked


def _parse_docx(source: Path, metadata: dict) -> NormalizedDocument:
    try:
        from docx import Document
    except ImportError:
        return _parse_loaded_units(source, metadata)
    document = Document(str(source))
    elements = []
    section_stack: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (getattr(paragraph.style, "name", "") or "").lower()
        if style_name.startswith("heading"):
            level = _heading_level(style_name)
            section_stack = section_stack[: max(level - 1, 0)] + [text]
            elements.append(
                _make_element(
                    source=source,
                    metadata=metadata,
                    index=len(elements),
                    element_type="heading",
                    content=text,
                    section_path=list(section_stack),
                    source_locator={"section_path": list(section_stack)},
                )
            )
            continue
        parent_id = _section_parent_id(metadata, section_stack)
        elements.append(
            _make_element(
                source=source,
                metadata=metadata,
                index=len(elements),
                element_type="paragraph",
                content=text,
                section_path=list(section_stack) or [metadata.get("topic") or source.stem],
                parent_id=parent_id,
                source_locator={"section_path": list(section_stack)},
            )
        )
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if not rows:
            continue
        section_path = list(section_stack) or [metadata.get("topic") or source.stem]
        table_element = _make_element(
            source=source,
            metadata=metadata,
            index=len(elements),
            element_type="table",
            content="\n".join(rows),
            section_path=section_path,
            source_locator={"table_index": table_index, "section_path": section_path},
            element_metadata={"table_index": table_index},
        )
        elements.append(table_element)
        for row_index, row_text in enumerate(rows, start=1):
            elements.append(
                _make_element(
                    source=source,
                    metadata=metadata,
                    index=len(elements),
                    element_type="table_row",
                    content=row_text,
                    section_path=section_path,
                    parent_id=table_element.element_id,
                    source_locator={"table_index": table_index, "row_index": row_index, "section_path": section_path},
                    element_metadata={"table_index": table_index, "row_index": row_index},
                )
            )
    return _base_document(source, metadata, elements)


def _parse_markdown(source: Path, metadata: dict) -> NormalizedDocument:
    text = source.read_text(encoding="utf-8", errors="ignore")
    elements = []
    section_stack: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        content = "\n".join(buffer).strip()
        if content:
            elements.append(
                _make_element(
                    source=source,
                    metadata=metadata,
                    index=len(elements),
                    element_type="list" if all(line.lstrip().startswith(("-", "*", "1.")) for line in buffer if line.strip()) else "paragraph",
                    content=content,
                    section_path=list(section_stack) or [source.stem],
                    parent_id=_section_parent_id(metadata, section_stack),
                    source_locator={"section_path": list(section_stack)},
                )
            )
        buffer = []

    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
        if match:
            flush()
            level = len(match.group(1))
            title = match.group(2).strip()
            section_stack = section_stack[: level - 1] + [title]
            elements.append(
                _make_element(
                    source=source,
                    metadata=metadata,
                    index=len(elements),
                    element_type="heading",
                    content=title,
                    section_path=list(section_stack),
                    source_locator={"section_path": list(section_stack)},
                )
            )
        elif line.strip():
            buffer.append(line)
        else:
            flush()
    flush()
    return _base_document(source, metadata, elements)


def _parse_xlsx(source: Path, metadata: dict) -> NormalizedDocument:
    try:
        import openpyxl
    except ImportError:
        return _parse_loaded_units(source, metadata)
    workbook = openpyxl.load_workbook(str(source), read_only=True, data_only=True)
    elements = []
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header_index = _find_header_index(rows)
            headers = [_clean_cell(value) for value in rows[header_index]]
            section_path = [sheet.title]
            table_parent = _make_element(
                source=source,
                metadata=metadata,
                index=len(elements),
                element_type="table",
                content=" | ".join(header for header in headers if header),
                section_path=section_path,
                source_locator={"sheet_name": sheet.title, "row_start": header_index + 1, "row_end": len(rows)},
                element_metadata={"sheet_name": sheet.title, "header_row": header_index + 1},
            )
            elements.append(table_parent)
            for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                values = [_clean_cell(value) for value in row]
                if not any(values):
                    continue
                pairs = [f"{header or 'column'}={value}" for header, value in zip(headers, values) if header or value]
                elements.append(
                    _make_element(
                        source=source,
                        metadata=metadata,
                        index=len(elements),
                        element_type="table_row",
                        content="; ".join(pairs),
                        section_path=section_path,
                        parent_id=table_parent.element_id,
                        source_locator={"sheet_name": sheet.title, "row_index": row_index},
                        element_metadata={"sheet_name": sheet.title, "row_index": row_index, "headers": headers},
                    )
                )
    finally:
        workbook.close()
    return _base_document(source, metadata, elements)


def _parse_csv(source: Path, metadata: dict) -> NormalizedDocument:
    with source.open("r", encoding="utf-8-sig", errors="ignore", newline="") as file:
        rows = list(csv.reader(file))
    elements = []
    if rows:
        header_index = _find_header_index(rows)
        headers = [_clean_cell(value) for value in rows[header_index]]
        table_parent = _make_element(
            source=source,
            metadata=metadata,
            index=0,
            element_type="table",
            content=" | ".join(header for header in headers if header),
            section_path=[source.stem],
            source_locator={"row_start": header_index + 1, "row_end": len(rows)},
            element_metadata={"header_row": header_index + 1},
        )
        elements.append(table_parent)
        for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            values = [_clean_cell(value) for value in row]
            if not any(values):
                continue
            pairs = [f"{header or 'column'}={value}" for header, value in zip(headers, values) if header or value]
            elements.append(
                _make_element(
                    source=source,
                    metadata=metadata,
                    index=len(elements),
                    element_type="table_row",
                    content="; ".join(pairs),
                    section_path=[source.stem],
                    parent_id=table_parent.element_id,
                    source_locator={"row_index": row_index},
                    element_metadata={"row_index": row_index, "headers": headers},
                )
            )
    return _base_document(source, metadata, elements)


def _parse_loaded_units(source: Path, metadata: dict) -> NormalizedDocument:
    elements = []
    for unit in load_source_units(source, metadata):
        locator = {}
        if unit.page_number:
            locator["page_number"] = unit.page_number
        if unit.sheet_name:
            locator["sheet_name"] = unit.sheet_name
        if unit.row_start:
            locator["row_start"] = unit.row_start
        if unit.row_end:
            locator["row_end"] = unit.row_end
        if unit.section:
            locator["section"] = unit.section
        element_type = "table" if unit.sheet_name or unit.metadata.get("table_index") else "paragraph"
        elements.append(
            _make_element(
                source=source,
                metadata=metadata,
                index=len(elements),
                element_type=element_type,
                content=unit.content,
                section_path=[unit.section or unit.title or source.stem],
                page_number=unit.page_number,
                source_locator=locator,
                element_metadata=dict(unit.metadata or {}),
            )
        )
    return _base_document(source, metadata, elements)


def _docling_requested() -> bool:
    return os.getenv("RAG_DOCLING_ENABLED", "false").strip().lower() in {"1", "true", "yes", "on"}


def _section_parent_id(metadata: dict, section_stack: list[str]) -> str | None:
    if not section_stack:
        return None
    return stable_element_id(
        document_id=stable_document_id(metadata["source_path"]),
        element_type="section",
        index=-1,
        content="/".join(section_stack),
        source_locator={"section_path": section_stack},
    )


def _heading_level(style_name: str) -> int:
    match = re.search(r"heading\s*(\d+)", style_name, re.I)
    return int(match.group(1)) if match else 1


def _find_header_index(rows: list[Any]) -> int:
    for index, row in enumerate(rows[:10]):
        values = [_clean_cell(value) for value in row]
        if sum(1 for value in values if value) >= 2:
            return index
    return 0


def _clean_cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalize_text(text: str) -> str:
    text = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines() if line.strip()).strip()
